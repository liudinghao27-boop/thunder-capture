"""Lead export generators (CSV/Excel)."""

import csv
import io
import json
import tempfile
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.utils import get_column_letter


DEFAULT_EXPORT_FIELDS = [
    "id",
    "platform",
    "keyword",
    "source_creator",
    "source_video_desc",
    "user_name",
    "unique_id",
    "short_id",
    "douyin_id",
    "text",
    "matched_categories",
    "ai_reply",
    "status",
    "error",
    "fetched_at",
    "processed_at",
]

FIELD_TITLES = {
    "id": "线索ID",
    "platform": "平台",
    "keyword": "来源关键词",
    "source_creator": "博主",
    "source_video_desc": "视频描述",
    "user_name": "用户昵称",
    "unique_id": "用户唯一ID",
    "short_id": "Short ID",
    "douyin_id": "抖音号",
    "text": "评论内容",
    "matched_categories": "意图分类",
    "ai_reply": "AI回复文案",
    "status": "状态",
    "error": "失败原因",
    "fetched_at": "采集时间",
    "processed_at": "处理时间",
}


def normalize_export_row(row: dict, fields: list[str]) -> dict:
    """Convert a TaskQueue dict/row into flat string values for export."""
    result = {}
    for field in fields:
        value = row.get(field)
        if value is None:
            value = ""
        elif isinstance(value, (list, tuple, set, dict)):
            value = json.dumps(value, ensure_ascii=False, default=str)
        else:
            value = str(value)
        result[field] = value
    return result


def _write_csv_rows(rows: Iterable[dict], fields: list[str]) -> Iterable[str]:
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=fields)
    writer.writerow({f: FIELD_TITLES.get(f, f) for f in fields})
    yield output.getvalue()
    for row in rows:
        output.seek(0)
        output.truncate(0)
        writer.writerow(normalize_export_row(row, fields))
        yield output.getvalue()


def generate_csv(rows: Iterable[dict], fields: list[str] | None = None) -> Iterable[str]:
    """Yield CSV lines as strings (StreamingResponse compatible)."""
    fields = fields or DEFAULT_EXPORT_FIELDS
    return _write_csv_rows(rows, fields)


def generate_xlsx(rows: Iterable[dict], fields: list[str] | None = None) -> Iterable[bytes]:
    """Yield an XLSX workbook as byte chunks.

    Uses an openpyxl write-only workbook and a temporary file so that the whole
    workbook is never materialised in memory. The input ``rows`` may be a
    generator, allowing the database to be read in batches.
    """
    fields = fields or DEFAULT_EXPORT_FIELDS
    headers = [FIELD_TITLES.get(f, f) for f in fields]

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        tmp_path = Path(tmp.name)

    try:
        wb = Workbook(write_only=True)
        ws = wb.create_sheet(title="线索池")
        ws.append(headers)

        for i, field in enumerate(fields, 1):
            col_letter = get_column_letter(i)
            ws.column_dimensions[col_letter].width = min(
                60, max(12, len(FIELD_TITLES.get(field, field)) + 2)
            )

        for row in rows:
            normalized = normalize_export_row(row, fields)
            ws.append([normalized[f] for f in fields])

        wb.save(tmp_path)

        with open(tmp_path, "rb") as f:
            while chunk := f.read(64 * 1024):
                yield chunk
    finally:
        tmp_path.unlink(missing_ok=True)
