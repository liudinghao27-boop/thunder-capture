import io
import csv
from openpyxl import load_workbook
from server.services.export import generate_csv, generate_xlsx, DEFAULT_EXPORT_FIELDS


def test_generate_csv_with_rows():
    rows = [
        {"id": 1, "platform": "douyin", "user_name": "u1", "text": "hello"},
        {"id": 2, "platform": "douyin", "user_name": "u2", "text": "world"},
    ]
    fields = ["id", "platform", "user_name", "text"]
    content = "".join(generate_csv(rows, fields=fields))
    reader = csv.reader(io.StringIO(content))
    data = list(reader)
    assert len(data) == 3  # header + 2 rows
    assert data[1][-1] == "hello"


def test_generate_csv_none_and_missing_fields():
    rows = [{"id": 1, "platform": None, "user_name": "u1"}]  # missing text
    fields = ["id", "platform", "user_name", "text"]
    content = "".join(generate_csv(rows, fields=fields))
    reader = csv.reader(io.StringIO(content))
    data = list(reader)
    assert data[1][1] == ""  # None -> empty
    assert data[1][3] == ""  # missing -> empty


def test_generate_csv_default_fields():
    rows = [{"id": 1}]
    content = "".join(generate_csv(rows))
    reader = csv.reader(io.StringIO(content))
    data = list(reader)
    assert len(data) == 2
    assert len(data[0]) == len(DEFAULT_EXPORT_FIELDS)


def test_generate_csv_empty_rows():
    content = "".join(generate_csv([]))
    reader = csv.reader(io.StringIO(content))
    data = list(reader)
    assert len(data) == 1
    assert len(data[0]) == len(DEFAULT_EXPORT_FIELDS)


def test_generate_xlsx_with_rows():
    rows = [
        {"id": 1, "platform": "douyin", "user_name": "u1", "text": "hello"},
    ]
    fields = ["id", "platform", "user_name", "text"]
    buffer = generate_xlsx(rows, fields=fields)
    wb = load_workbook(buffer)
    ws = wb.active
    assert ws.title == "线索池"
    assert ws.cell(row=1, column=1).value == "线索ID"
    assert ws.cell(row=2, column=2).value == "douyin"
    assert ws.cell(row=2, column=4).value == "hello"


def test_generate_xlsx_collection_field():
    rows = [{"id": 1, "matched_categories": ["咨询", "其他"]}]
    fields = ["id", "matched_categories"]
    buffer = generate_xlsx(rows, fields=fields)
    wb = load_workbook(buffer)
    ws = wb.active
    assert "咨询" in ws.cell(row=2, column=2).value
